"""Build the one-page resume with all six linked projects.

Design preset: standard_business_brief, with the named ats_resume override.
Letter, 0.55in top/bottom, 0.62in sides; Arial, 10pt body / 11.6pt leading;
28pt name; 9.5pt section headings; restrained teal accents and section rules.
Single main reading column, a three-row skills table, native paragraphs/bullets/hyperlinks,
and no text boxes, images, or contact details in headers/footers. PDF embeds Arial.
"""
import json
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, BaseDocTemplate, Frame, PageTemplate, Table, TableStyle, Flowable, KeepTogether

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'career'
OUT.mkdir(exist_ok=True)
data = json.loads((ROOT / 'content-review/career-content.json').read_text())
profile = data['profile']
INK, ACCENT, MUTED, RULE = '172B40', '24647B', '596878', 'C9D5DE'
TINT = 'EDF3F7'
CONTENT_WIDTH = 612 - 2 * 44.64
MARGIN_X, MARGIN_Y = 44.64, 39.6
FONT_DIR = Path('/System/Library/Fonts/Supplemental')
for name, file in [('ResumeArial', 'Arial.ttf'), ('ResumeArialBold', 'Arial Bold.ttf')]:
    path = FONT_DIR / file
    if not path.exists():
        raise FileNotFoundError(f'Arial font required for matching PDF/Word metrics: {path}')
    pdfmetrics.registerFont(TTFont(name, str(path)))
pdfmetrics.registerFontFamily('ResumeArial', normal='ResumeArial', bold='ResumeArialBold')

# Compact descriptions summarize only reviewed facts. The website retains the detail.
experience_summaries = {
    'Microble Technologies': 'Develop industrial inspection pipelines for insulator defect detection with YOLO and OpenCV; own dataset preparation, annotation, augmentation, model training, evaluation, and deployment.',
    'Empire Circuits LLC': 'Built PondGuard’s detection and response software, an OpenCV PCB image-stitching pipeline, and ThingsBoard/Grafana monitoring dashboards backed by MQTT.',
    'Infotact Solution': 'Built Droplify with Flask, SQLite, Beautiful Soup, and Selenium, including product-data collection, analytics, and automated price-drop alerts.',
    'CHARUSAT University': 'Led a four-engineer backend team on Placestar: Node.js/Express APIs, MySQL schemas, JWT access controls, sprints, and reviews. Supported a live examination with 100+ students.',
}
project_summaries = {
    'PondGuard': ('Python · Flask · YOLO11', 'Owned bird-detection software with event capture and automated sprinkler and red-beam deterrents.'),
    'Placestar': ('Node.js · Express · MySQL', 'Built placement/exam APIs with Wi-Fi restrictions and single-session controls; supported 100+ students.'),
    'Droplify': ('Flask · SQLite · Selenium', 'Solo-built four-platform price tracking, analytics, and email alerts; validated 25+ product pages.'),
    'AMC Connect': ('FastAPI · PyTorch · YOLO', 'Built civic-complaint APIs, GPS validation, image screening, and LLM-generated descriptions.'),
    'ClubSphere': ('Node.js · PostgreSQL · JWT', 'Led club-management APIs, database design, and authentication; piloted at CHARUSAT University.'),
    'Worker Location Management': ('ThingsBoard · Geofencing', 'Configured zone/inactivity rules and alerts; validated monitoring behavior using RSSI signal data.'),
}
assert set(project_summaries) == {p['title'] for p in data['projects']}

def link(label, url):
    return f'<a href="{escape(url, {chr(34): "&quot;"})}">{escape(label)}</a>'

def muted(text):
    return f'<font color="#{MUTED}">{escape(text)}</font>'

blocks = []
def add(role, markup):
    blocks.append((role, markup))

add('Title', escape(profile['name']))
add('Subtitle', 'SOFTWARE DEVELOPER  |  PYTHON, BACKEND &amp; COMPUTER VISION')
add('Contact', f"{escape(profile['location'])}  |  {link(profile['email'], 'mailto:' + profile['email'])}  |  {escape(profile['phone'])}")
add('Contact', link(profile['website'].removeprefix('https://'), profile['website']) + '  |  ' + link(profile['github'].removeprefix('https://'), profile['github']))
add('Contact', link(profile['linkedin'].removeprefix('https://').rstrip('/'), profile['linkedin']))
add('Heading 1', 'PROFESSIONAL SUMMARY')
add('Normal', 'Software developer focused on Python backends and computer vision, with experience in industrial inspection, REST API architecture, and backend team leadership.')
add('Heading 1', 'TECHNICAL SKILLS')
add('Skill Row', '<b>Languages &amp; backend:</b> Python, JavaScript, SQL, Flask, FastAPI, Django, Node.js, Express, REST APIs, JWT')
add('Skill Row', '<b>Vision &amp; data:</b> YOLO, OpenCV, PyTorch, PostgreSQL, MySQL, SQLite, MongoDB, Supabase')
add('Skill Row', '<b>Tools:</b> Git, Docker, Linux, Azure DevOps, CI/CD, Postman, Beautiful Soup, Selenium')
add('Heading 1', 'PROFESSIONAL EXPERIENCE')
for item in data['experience']:
    dates = item['year'].replace(' — ', '–')
    add('Heading 2', f"<b>{escape(item['title'])}</b> | {escape(item['company'])}  <date>{escape(dates)}</date>")
    add('List Bullet', escape(experience_summaries[item['company']]))
add('Heading 1', 'PROJECTS')
project_urls = []
for item in data['projects']:
    # Match the actual article IDs in src/components/portfolio/Projects.tsx.
    url = profile['website'].rstrip('/') + '/#project-' + item['title'].lower().replace(' ', '-')
    project_urls.append(url)
    stack, description = project_summaries[item['title']]
    add('Project Heading', f"<b>{link(item['title'], url)}</b>  {muted('| ' + stack)}")
    add('Project Body', escape(description))
add('Heading 1', 'EDUCATION')
for item in data['education']:
    add('Heading 2', f"<b>{escape(item['degree'])}</b>  <date>{escape(item['year'].replace(' — ', '–'))}</date>")
    add('Education Body', escape(item['institution'] + ' | ' + item['detail'].replace(' / 10.00', '/10')))

# Font size, leading, before, after, color. Explicit tokens for every text style.
spec = {
    'Title': (28, 31, 0, 4, INK),
    'Subtitle': (9.5, 12, 0, 6, ACCENT),
    'Contact': (9, 11, 0, 1, MUTED),
    'Normal': (10, 11.6, 0, 2, INK),
    'Skill Row': (9.5, 11.4, 0, 0, INK),
    'Heading 1': (9.5, 12, 9, 5, INK),
    'Heading 2': (10, 12, 4, 2, INK),
    'Heading 3': (10, 12, 4, 2, INK),
    'List Bullet': (10, 11.6, 0, 2, INK),
    'Project Heading': (9.8, 11.7, 0, 1, INK),
    'Project Body': (9.5, 11.2, 0, 0, INK),
    'Education Body': (9.5, 11.4, 0, 1, MUTED),
}
heading_roles = {'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Project Heading'}

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Pt(MARGIN_Y)
section.left_margin = section.right_margin = Pt(MARGIN_X)
section.header_distance = section.footer_distance = Inches(.492)
for role, (size, leading, before, after, color) in spec.items():
    style = doc.styles[role] if role in doc.styles else doc.styles.add_style(role, 1)
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style.font.bold = role in {'Title', 'Subtitle', 'Heading 1'}
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = style.paragraph_format
    fmt.line_spacing = Pt(leading)
    fmt.space_before, fmt.space_after = Pt(before), Pt(after)
    fmt.keep_with_next = role in heading_roles
    fmt.keep_together = True
    fmt.widow_control = True
    # Clear template borders before adding the explicit section rule.
    ppr = style.element.get_or_add_pPr()
    for border in list(ppr.findall(qn('w:pBdr'))):
        ppr.remove(border)
    if role in {'Heading 2'}:
        fmt.tab_stops.add_tab_stop(Pt(CONTENT_WIDTH), WD_TAB_ALIGNMENT.RIGHT)
    if role == 'Heading 1':
        shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), TINT); ppr.append(shading)
        borders = OxmlElement('w:pBdr')
        line = OxmlElement('w:bottom')
        for key, value in {'val': 'single', 'sz': '6', 'space': '3', 'color': RULE}.items():
            line.set(qn('w:' + key), value)
        borders.append(line)
        ppr.append(borders)
# Define a real list with a small hanging indent (no typed bullet characters).
numbering = doc.part.numbering_part.element
abstract_id = max(int(n.get(qn('w:abstractNumId'))) for n in numbering.findall(qn('w:abstractNum'))) + 1
num_id = max(int(n.get(qn('w:numId'))) for n in numbering.findall(qn('w:num'))) + 1
abstract = OxmlElement('w:abstractNum'); abstract.set(qn('w:abstractNumId'), str(abstract_id))
lvl = OxmlElement('w:lvl'); lvl.set(qn('w:ilvl'), '0')
for tag, value in [('start', '1'), ('numFmt', 'bullet'), ('lvlText', '•'), ('lvlJc', 'left')]:
    element = OxmlElement('w:' + tag); element.set(qn('w:val'), value); lvl.append(element)
ppr = OxmlElement('w:pPr'); ind = OxmlElement('w:ind')
ind.set(qn('w:left'), '200'); ind.set(qn('w:hanging'), '150'); ppr.append(ind)
tabs = OxmlElement('w:tabs'); tab = OxmlElement('w:tab')
tab.set(qn('w:val'), 'num'); tab.set(qn('w:pos'), '200'); tabs.append(tab); ppr.append(tabs); lvl.append(ppr)
abstract.append(lvl); numbering.append(abstract)
num = OxmlElement('w:num'); num.set(qn('w:numId'), str(num_id))
ref = OxmlElement('w:abstractNumId'); ref.set(qn('w:val'), str(abstract_id)); num.append(ref); numbering.append(num)

# Render the same small markup vocabulary as native Word runs and hyperlinks.
def add_docx_markup(paragraph, markup, role):
    base = spec[role]
    def emit(text, bold=False, color=None, href=None):
        if not text:
            return
        run = paragraph.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(base[0])
        run.font.bold = bold or role in {'Title', 'Subtitle', 'Heading 1'}
        run.font.color.rgb = RGBColor.from_string(color or base[4])
        if href:
            run.font.underline = role == 'Project Heading'
            run.font.color.rgb = RGBColor.from_string(ACCENT if role == 'Project Heading' else base[4])
            hyperlink = OxmlElement('w:hyperlink')
            hyperlink.set(qn('r:id'), doc.part.relate_to(href, RT.HYPERLINK, is_external=True))
            hyperlink.append(run._r)
            paragraph._p.append(hyperlink)
    def walk(node, bold=False, color=None, href=None):
        if node.tag == 'date':
            emit('\t' + (node.text or ''), False, MUTED)
            return
        bold = bold or node.tag == 'b'
        color = node.attrib.get('color', color)
        if color:
            color = color.lstrip('#')
        href = node.attrib.get('href', href)
        emit(node.text, bold, color, href)
        for child in node:
            walk(child, bold, color, href)
            emit(child.tail, bold, color, href)
    walk(ET.fromstring('<root>' + markup + '</root>'))

skill_table = None
project_index = -1
for role, markup in blocks:
    if role == 'Skill Row':
        if skill_table is None:
            skill_table = doc.add_table(rows=0, cols=2)
            skill_table.alignment = WD_TABLE_ALIGNMENT.LEFT
            skill_table.autofit = False
            widths = [2000, 8454]
            props = skill_table._tbl.tblPr
            props.find(qn('w:tblW')).set(qn('w:w'), str(sum(widths)))
            props.find(qn('w:tblW')).set(qn('w:type'), 'dxa')
            indent = OxmlElement('w:tblInd'); indent.set(qn('w:w'), '90'); indent.set(qn('w:type'), 'dxa'); props.append(indent)
            margins = OxmlElement('w:tblCellMar')
            for side, value in [('top', 55), ('bottom', 55), ('start', 90), ('end', 90)]:
                el = OxmlElement('w:' + side); el.set(qn('w:w'), str(value)); el.set(qn('w:type'), 'dxa'); margins.append(el)
            props.append(margins)
            borders = OxmlElement('w:tblBorders')
            for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
                el = OxmlElement('w:' + side)
                for key, value in {'val': 'single', 'sz': '4', 'color': RULE}.items(): el.set(qn('w:' + key), value)
                borders.append(el)
            props.append(borders)
            for grid, width in zip(skill_table._tbl.tblGrid, widths): grid.set(qn('w:w'), str(width))
        label, values = markup.split('</b> ', 1)
        cells = skill_table.add_row().cells
        for cell, width, text in zip(cells, [2000, 8454], [label + '</b>', values]):
            cell.width = Pt(width / 20)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]; paragraph.style = doc.styles['Skill Row']
            add_docx_markup(paragraph, text, 'Skill Row')
        shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), TINT); cells[0]._tc.get_or_add_tcPr().append(shading)
        continue
    paragraph = doc.add_paragraph(style=role)
    add_docx_markup(paragraph, markup, role)
    if role in {'Project Heading', 'Project Body'}:
        if role == 'Project Heading': project_index += 1
        props = paragraph._p.get_or_add_pPr()
        borders = OxmlElement('w:pBdr')
        for side in ['left', 'right', 'top' if role == 'Project Heading' else 'bottom']:
            line = OxmlElement('w:' + side)
            for key, value in {'val': 'single', 'sz': '4', 'space': '4', 'color': RULE}.items(): line.set(qn('w:' + key), value)
            borders.append(line)
        props.append(borders)
        shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), 'F4F7FA' if project_index % 2 == 0 else 'FFFFFF'); props.append(shading)
        paragraph.paragraph_format.left_indent = Pt(6)
        paragraph.paragraph_format.right_indent = Pt(6)
        if role == 'Project Heading': paragraph.paragraph_format.space_before = Pt(4)
        else: paragraph.paragraph_format.space_after = Pt(4)
    if role == 'List Bullet':
        props = paragraph._p.get_or_add_pPr()
        numpr = OxmlElement('w:numPr')
        level = OxmlElement('w:ilvl'); level.set(qn('w:val'), '0'); numpr.append(level)
        numref = OxmlElement('w:numId'); numref.set(qn('w:val'), str(num_id)); numpr.append(numref); props.append(numpr)
doc.core_properties.title = 'Krish Lalani | Software Developer'
doc.core_properties.author = 'Krish Lalani'
doc.core_properties.subject = 'Python, backend engineering and computer vision'
doc.core_properties.created = doc.core_properties.modified = datetime.now(timezone.utc)
doc.save(OUT / 'Krish_Lalani_Resume.docx')

class SectionHeading(Paragraph):
    """Native text with a quiet section rule, preserving the text reading order."""
    def draw(self):
        self.canv.setFillColor(colors.HexColor('#' + TINT))
        self.canv.rect(0, -2, self.width, self.height + 5, fill=1, stroke=0)
        super().draw()
        self.canv.setStrokeColor(colors.HexColor('#' + RULE))
        self.canv.setLineWidth(.5)
        self.canv.line(0, -2, self.width, -2)

pdf_styles = {}
for role, (size, leading, before, after, color) in spec.items():
    pdf_styles[role] = ParagraphStyle(
        role, fontName='ResumeArialBold' if role in {'Title', 'Subtitle', 'Heading 1'} else 'ResumeArial',
        fontSize=size, leading=leading, spaceBefore=before, spaceAfter=after,
        textColor=colors.HexColor('#' + color), keepWithNext=role in heading_roles,
        leftIndent=10 if role == 'List Bullet' else 0, bulletIndent=2.5,
        bulletFontName='ResumeArial', bulletFontSize=8,
    )
class DatedHeading(Flowable):
    def __init__(self, markup, style):
        super().__init__()
        left, right = markup.split('<date>')
        self.left = Paragraph(left.rstrip(), style)
        self.right = Paragraph(right.removesuffix('</date>'), ParagraphStyle('Date', parent=style, alignment=2, textColor=colors.HexColor('#'+MUTED), fontSize=9))
        self.spaceBefore, self.spaceAfter, self.keepWithNext = 4, 2, True
    def wrap(self, width, height):
        self.width = width
        _, lh = self.left.wrap(width - 110, height)
        _, rh = self.right.wrap(105, height)
        self.height = max(lh, rh)
        return width, self.height
    def draw(self):
        self.left.drawOn(self.canv, 0, self.height-self.left.height)
        self.right.drawOn(self.canv, self.width-105, self.height-self.right.height)

class ProjectRow(Flowable):
    def __init__(self, heading, description, index):
        super().__init__()
        self.heading = Paragraph(heading, pdf_styles['Project Heading'])
        self.description = Paragraph(description, pdf_styles['Project Body'])
        self.index = index
    def wrap(self, width, height):
        self.width = width
        self.heading.wrap(width-14, height)
        self.description.wrap(width-14, height)
        self.height = self.heading.height+self.description.height+9
        return self.width, self.height
    def draw(self):
        self.canv.setFillColor(colors.HexColor('#F4F7FA' if self.index%2==0 else '#FFFFFF'))
        self.canv.setStrokeColor(colors.HexColor('#'+RULE))
        self.canv.setLineWidth(.4)
        self.canv.rect(0, 0, self.width, self.height, fill=1, stroke=1)
        self.heading.drawOn(self.canv, 7, self.height-4-self.heading.height)
        self.description.drawOn(self.canv, 7, 4)

story = []
index = 0
project_index = 0
while index < len(blocks):
    role, markup = blocks[index]
    if role == 'Skill Row':
        rows = []
        while index < len(blocks) and blocks[index][0] == 'Skill Row':
            label, values = blocks[index][1].split('</b> ', 1)
            rows.append([Paragraph(label+'</b>', pdf_styles['Skill Row']), Paragraph(values, pdf_styles['Skill Row'])])
            index += 1
        table = Table(rows, colWidths=[100, CONTENT_WIDTH-100], hAlign='LEFT')
        table.setStyle(TableStyle([
            ('GRID', (0,0), (-1,-1), .4, colors.HexColor('#'+RULE)),
            ('BACKGROUND', (0,0), (0,-1), colors.HexColor('#'+TINT)),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('LEFTPADDING', (0,0), (-1,-1), 4.5), ('RIGHTPADDING', (0,0), (-1,-1), 4.5),
            ('TOPPADDING', (0,0), (-1,-1), 2.75), ('BOTTOMPADDING', (0,0), (-1,-1), 2.75),
        ]))
        story.append(table)
        continue
    if role == 'Project Heading':
        markup = markup.replace('<a ', f'<font color="#{ACCENT}"><u><a ').replace('</a>', '</a></u></font>')
        story.append(ProjectRow(markup, blocks[index+1][1], project_index))
        project_index += 1
        index += 2
        continue
    if '<date>' in markup:
        story.append(DatedHeading(markup, pdf_styles[role]))
    else:
        cls = SectionHeading if role == 'Heading 1' else Paragraph
        story.append(cls(markup, pdf_styles[role], bulletText='•' if role == 'List Bullet' else None))
    index += 1

def page_decoration(canvas, document):
    canvas.setStrokeColor(colors.HexColor('#D8E1E8'))
    canvas.setLineWidth(.6)
    canvas.rect(22, 22, 568, 748, fill=0, stroke=1)
    canvas.setFillColor(colors.HexColor('#'+ACCENT))
    canvas.rect(22, 735, 3, 35, fill=1, stroke=0)

pdf = BaseDocTemplate(str(OUT / 'Krish_Lalani_Resume.pdf'), pagesize=(612,792), title='Krish Lalani | Software Developer', author='Krish Lalani')
pdf.addPageTemplates(PageTemplate(id='Resume', frames=[Frame(MARGIN_X, MARGIN_Y, CONTENT_WIDTH, 792-2*MARGIN_Y, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)], onPage=page_decoration))
pdf.build(story)
reader = PdfReader(OUT / 'Krish_Lalani_Resume.pdf')
text = '\n'.join(page.extract_text() for page in reader.pages)
(OUT / 'Krish_Lalani_Resume.txt').write_text(text)
assert len(reader.pages) == 1, f'Resume must be one page; got {len(reader.pages)}'
assert len(doc.tables) == 1 and len(doc.tables[0].rows) == 3
for value in [profile['name'], profile['phone'], 'PROFESSIONAL EXPERIENCE', 'TECHNICAL SKILLS', 'PROJECTS', 'EDUCATION', '100+', '25+', '9.07', '9.89']:
    assert value in text, value
pdf_links = [annotation.get_object().get('/A', {}).get('/URI') for annotation in reader.pages[0].get('/Annots', [])]
docx_links = [rel.target_ref for rel in doc.part.rels.values() if rel.reltype == RT.HYPERLINK]
for project, url in zip(data['projects'], project_urls):
    assert project['title'] in text
    assert url in pdf_links, url
    assert url in docx_links, url
print(f'Built PDF (verified 1 page) and Word resume: {len(text.split())} words; all 6 projects and their PDF/Word links verified.')
